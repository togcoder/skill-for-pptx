#!/usr/bin/env python3
"""Unit tests for T024 complex motion upgrades:
- Safe non-root environment probe (check_environment.py)
- Multi-segment poly-path compilation (poly_path / via keyframes)
- Color (p:animClr), Teeter, repeatCount, autoRev, p14:bounceEnd
- Interactive shape-triggered sequences (p:interactiveSeq / trigger_shape)
- Group-aware virtual rigging for separate shapes + overlay text boxes
- Mid-flight trajectory collision detection & auto-arc obstacle avoidance
- New orbit and wave recipes
- Morph empty-placeholder detection & stripping
"""
import json
import math
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from lxml import etree as E
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"scripts"))

import motion_director as md
import motion_engine as me
import pptx_animator as anim

SHOWCASE=ROOT/"tests"/"fixtures"/"motion_showcase_deck.pptx"
NS=anim.NS


def slide_root(path,part):
    with ZipFile(path) as z:
        return E.fromstring(z.read(part))


def build_rigged_and_obstacle_deck(path):
    prs=Presentation()
    prs.slide_width=Inches(13.333)
    prs.slide_height=Inches(7.5)
    title_layout,title_only=prs.slide_layouts[0],prs.slide_layouts[5]

    s1=prs.slides.add_slide(title_layout)
    s1.shapes.title.text="T024 Complex Motion Test Deck"
    s1.placeholders[1].text="Synthetic fixture for virtual rigging, auto-arc and interactiveSeq"

    s2=prs.slides.add_slide(title_only)
    s2.shapes.title.text="Cycle with Separate Labels (Virtual Rigging)"
    cx,cy,r=6.67,4.3,2.0
    hub=s2.shapes.add_shape(MSO_SHAPE.OVAL,Inches(cx-1.0),Inches(cy-1.0),Inches(2.0),Inches(2.0))
    hub.name="Core Hub"
    hub_lbl=s2.shapes.add_textbox(Inches(cx-0.8),Inches(cy-0.3),Inches(1.6),Inches(0.6))
    hub_lbl.name="Core Hub Label"
    hub_lbl.text_frame.text="AI Engine"
    for i,label in enumerate(["Sense","Plan","Act","Learn"]):
        ang=-math.pi/2+i*math.pi/2
        x,y=cx+r*math.cos(ang),cy+r*math.sin(ang)
        node=s2.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x-0.85),Inches(y-0.85),Inches(1.7),Inches(1.7))
        node.name=f"Node {label}"
        node.fill.solid()
        node.fill.fore_color.rgb=RGBColor(0x0F,0x9D,0x8A)
        tb=s2.shapes.add_textbox(Inches(x-0.65),Inches(y-0.25),Inches(1.3),Inches(0.5))
        tb.name=f"Label {label}"
        tb.text_frame.text=label

    s3=prs.slides.add_slide(title_only)
    s3.shapes.title.text="Obstacle Avoidance & Interactive Trigger"
    tok=s3.shapes.add_shape(MSO_SHAPE.OVAL,Inches(1.2),Inches(3.4),Inches(0.7),Inches(0.7))
    tok.name="Packet"
    tok.text_frame.text="P1"
    wall=s3.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(5.8),Inches(2.6),Inches(1.7),Inches(1.8))
    wall.name="Firewall Blocker"
    wall.text_frame.text="Firewall"
    dest=s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(10.2),Inches(3.8),Inches(1.8),Inches(1.2))
    dest.name="Gateway"
    dest.text_frame.text="Gateway"
    btn=s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(5.4),Inches(6.0),Inches(2.5),Inches(0.8))
    btn.name="Inspect Button"
    btn.text_frame.text="Click to Inspect"
    detail=s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(8.3),Inches(5.8),Inches(3.8),Inches(1.1))
    detail.name="Detail Panel"
    detail.text_frame.text="TLS 1.3 Verified · 0 Drops"

    prs.save(path)


class T024ComplexMotionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory()
        cls.td=Path(cls.tmp.name)
        cls.deck=cls.td/"rigged_obstacle.pptx"
        build_rigged_and_obstacle_deck(cls.deck)
        cls.model=md.deck_model(cls.deck)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_check_environment_runs_without_permission_error(self):
        res=subprocess.run([sys.executable,str(ROOT/"scripts"/"check_environment.py")],
                           capture_output=True,text=True,check=True)
        data=json.loads(res.stdout)
        self.assertTrue(data["lxml_available"])
        self.assertIn("pillow_available",data)

    def test_poly_segment_path_compiles_to_single_anim_motion(self):
        s3=self.model["slides"][2]
        by={o["name"]:o for o in s3["objects"]}
        st={o["id"]:me.fresh_state() for o in s3["objects"]}
        pkt=by["Packet"]
        effs=me.compile_track({
            "poly_path":True,
            "keyframes":[
                {"t":0},
                {"t":400,"dx":0.2,"dy":-0.15,"curve":0.2},
                {"t":800,"dx":0.4,"dy":0.0,"curve":-0.2},
                {"t":1200,"dx":0.6,"dy":-0.1},
            ]
        },pkt,by,st[pkt["id"]],"poly")
        paths=[e for e in effs if e["preset"]=="path"]
        self.assertEqual(len(paths),1)
        self.assertEqual(len(paths[0]["anchored"]),4)
        self.assertEqual(paths[0]["duration_ms"],1200)
        mid=me.state_at(st,effs,600)[pkt["id"]]
        self.assertGreater(mid["dx"],0.15)
        self.assertLess(mid["dx"],0.45)

    def test_color_teeter_repeat_autorev_and_bounce_round_trip(self):
        root=slide_root(SHOWCASE,"ppt/slides/slide5.xml")
        anim.apply_timeline(root,[{"id":"c1","effects":[
            {"preset":"path","spid":"7","trigger":"click","duration_ms":800,
             "anchored":[{"x":0,"y":0},{"x":0.2,"y":0}],"bounce_end":0.45,"loop":{"repeat":2,"auto_reverse":True}},
            {"preset":"color","spid":"3","trigger":"with","duration_ms":500,"color":"#E06C4F"},
            {"preset":"teeter","spid":"4","trigger":"with","duration_ms":600,"by_deg":3.0},
        ]}])
        groups,_=me.effects_from_slide(root)
        self.assertEqual(len(groups),1)
        by_preset={e["preset"]:e for e in groups[0]}
        self.assertAlmostEqual(by_preset["path"]["bounce_end"],0.45,places=3)
        self.assertEqual(by_preset["path"]["loop"],{"repeat":2,"auto_reverse":True})  # T032: one loop model
        self.assertEqual(by_preset["color"]["color"],"#E06C4F")
        self.assertEqual(by_preset["teeter"]["preset"],"teeter")
        # T032: written with PowerPoint's own presets (Teeter 32, Change Fill Color 1/2), not Blast/Line Color.
        ids={(c.get("presetClass"),c.get("presetID"),c.get("presetSubtype")) for c in root.iter(anim.q("cTn")) if c.get("presetClass")}
        self.assertIn(("emph","32","0"),ids)
        self.assertIn(("emph","1","2"),ids)

    def test_interactive_sequence_writes_and_reads_back(self):
        root=slide_root(self.deck,"ppt/slides/slide3.xml")
        by={o["name"]:o["id"] for o in self.model["slides"][2]["objects"]}
        receipt=anim.apply_timeline(root,[
            {"id":"main-1","effects":[{"preset":"fade","spid":by["Packet"],"trigger":"click","duration_ms":400}]},
            {"id":"trig-1","trigger_spid":by["Inspect Button"],"effects":[
                {"preset":"zoom","spid":by["Detail Panel"],"trigger":"click","duration_ms":500}
            ]},
        ])
        self.assertEqual(receipt["appended_click_groups"],1)
        self.assertEqual(receipt["interactive_click_groups"],1)
        inter=anim.read_interactive_sequences(root)
        self.assertEqual(len(inter),1)
        self.assertEqual(inter[0]["trigger_spid"],by["Inspect Button"])
        self.assertEqual(inter[0]["effects"][0]["spid"],by["Detail Panel"])

    def test_virtual_rigging_detects_separate_label_cycle_and_moves_followers_in_sync(self):
        s2=self.model["slides"][1]
        rigs=me.detect_virtual_rigs(s2["objects"])
        by_name={o["name"]:o for o in s2["objects"]}
        self.assertIn(by_name["Node Sense"]["id"],rigs)
        follower_names=[f["name"] for f in rigs[by_name["Node Sense"]["id"]]]
        self.assertEqual(follower_names,["Label Sense"])

        plan=md.draft(self.model,style="cinematic")
        s2_plan=next(s for s in plan["slides"] if s["source_index"]==2)
        recipes=[b.get("recipe") for b in s2_plan["beats"]]
        self.assertIn("assemble",recipes)
        self.assertIn("spotlight",recipes)

        nodes=[by_name[f"Node {k}"] for k in ("Sense","Plan","Act","Learn")]
        tok={o["name"]:o for o in s2["objects"]}
        st={o["id"]:me.fresh_state() for o in s2["objects"]}
        effs=me.compile_choreography("cycle",nodes,st,tok,"cyc",{"duration_ms":800})
        end=me.advance(st,effs)
        node_sense=by_name["Node Sense"]
        lbl_sense=by_name["Label Sense"]
        node_plan=by_name["Node Plan"]
        self.assertAlmostEqual(end[node_sense["id"]]["dx"],end[lbl_sense["id"]]["dx"],places=6)
        self.assertAlmostEqual(end[node_sense["id"]]["dy"],end[lbl_sense["id"]]["dy"],places=6)
        self.assertAlmostEqual(me.authored(node_sense)["cx"]+end[node_sense["id"]]["dx"],
                               me.authored(node_plan)["cx"],places=5)

    def test_mid_flight_collision_detection_and_auto_arc_avoidance(self):
        s3=self.model["slides"][2]
        by={o["name"]:o for o in s3["objects"]}
        st={o["id"]:me.fresh_state() for o in s3["objects"]}
        straight=me.compile_choreography("travel",[by["Packet"],by["Gateway"]],st,by,"tr1",
                                         {"arc":0.0,"offset_y":-0.14,"auto_avoid":False,"pulse_stops":False})
        end_st=me.advance(st,straight)
        self.assertEqual(me.layout_warnings(s3["objects"],end_st,"end"),[])
        trans_w=me.transient_warnings(s3["objects"],st,straight,"mid")
        self.assertTrue(any("Firewall Blocker" in w for w in trans_w),f"Expected mid-motion warning, got {trans_w}")

        curved=me.compile_choreography("travel",[by["Packet"],by["Gateway"]],st,by,"tr2",
                                       {"arc":0.0,"offset_y":-0.14,"auto_avoid":True,"pulse_stops":False})
        self.assertIn("c1",curved[0]["anchored"][-1])
        self.assertEqual(me.transient_warnings(s3["objects"],st,curved,"mid"),[])

    def test_orbit_and_wave_recipes(self):
        s2=self.model["slides"][1]
        by={o["name"]:o for o in s2["objects"]}
        nodes=[by["Core Hub"]]+[by[f"Node {k}"] for k in ("Sense","Plan","Act","Learn")]
        st={o["id"]:me.fresh_state() for o in s2["objects"]}
        orbit_effs=me.compile_choreography("orbit",nodes,st,by,"orb",
                                           {"around_first":True,"angle_deg":360,"duration_ms":1200})
        sense_path=[e for e in orbit_effs if e["spid"]==by["Node Sense"]["id"] and e["preset"]=="path"][0]
        self.assertEqual(len(sense_path["anchored"]),5)
        end_st=me.advance(st,orbit_effs)
        self.assertAlmostEqual(end_st[by["Node Sense"]["id"]]["dx"],0.0,places=5)
        self.assertAlmostEqual(end_st[by["Node Sense"]["id"]]["dy"],0.0,places=5)

        wave_effs=me.compile_choreography("wave",nodes[1:],st,by,"wv",{"duration_ms":600,"stagger_ms":100})
        end_wave=me.advance(st,wave_effs)
        for n in nodes[1:]:
            self.assertAlmostEqual(end_wave[n["id"]]["dy"],0.0,places=5)
            self.assertAlmostEqual(end_wave[n["id"]]["scale"],1.0,places=5)

    def test_empty_placeholders_detected_and_stripped_for_morph(self):
        prs=Presentation()
        slide=prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text="Morph Target With Empty Body Placeholder"
        p=self.td/"empty_ph.pptx"
        prs.save(p)
        root=slide_root(p,"ppt/slides/slide1.xml")
        empty=anim.empty_placeholders(root)
        self.assertEqual(len(empty),1)
        anim.add_morph_transition(root,1200,strip_empty_placeholders=True)
        self.assertEqual(anim.empty_placeholders(root),[])
        self.assertIsNotNone(root.find(f".//{{{anim.P159}}}morph"))


    def test_follow_through_orbit3d_explode_and_cascade_flow(self):
        s2=self.model["slides"][1]
        by={o["name"]:o for o in s2["objects"]}
        nodes=[by["Core Hub"]]+[by[f"Node {k}"] for k in ("Sense","Plan","Act","Learn")]
        st={o["id"]:me.fresh_state() for o in s2["objects"]}

        # 1. follow_through_ms lags follower behind carrier
        asm=me.compile_choreography("assemble",nodes[1:],st,by,"asm",
                                    {"from":"center","duration_ms":600,"stagger_ms":100,"follow_through_ms":40})
        node_moves=[e for e in asm if e["spid"]==by["Node Sense"]["id"] and e["preset"]=="path"]
        lbl_moves=[e for e in asm if e["spid"]==by["Label Sense"]["id"] and e["preset"]=="path"]
        self.assertEqual(lbl_moves[0]["delay_ms"]-node_moves[0]["delay_ms"],40)
        self.assertEqual(lbl_moves[0]["duration_ms"],node_moves[0]["duration_ms"])

        # 2. orbit-3d modulates scale and opacity by elliptical depth
        o3d=me.compile_choreography("orbit-3d",nodes,st,by,"o3d",
                                    {"around_first":True,"steps":1,"duration_ms":900,"depth_scale":0.22})
        end_o3d=me.advance(st,o3d)
        # Node Plan (east, moves to south/front where sin(pi/2)=+1) grows > 1.15
        self.assertGreater(end_o3d[by["Node Plan"]["id"]]["scale"],1.15)
        # Node Learn (west, moves to north/back where sin(-pi/2)=-1) shrinks < 0.85
        self.assertLess(end_o3d[by["Node Learn"]["id"]]["scale"],0.85)
        # Follower Label Plan scales in lockstep with Node Plan
        self.assertAlmostEqual(end_o3d[by["Node Plan"]["id"]]["scale"],
                               end_o3d[by["Label Plan"]["id"]]["scale"],places=4)

        # 3. explode-layers separates stacked strata along isometric axis
        exp=me.compile_choreography("explode-layers",nodes[1:],st,by,"exp",
                                    {"axis":"iso","spread":0.24,"stagger_ms":60,"bounce_end":0.3})
        end_exp=me.advance(st,exp)
        dys=[end_exp[n["id"]]["dy"] for n in nodes[1:]]
        self.assertLess(dys[0],0)
        self.assertGreater(dys[-1],0)

        # 4. cascade-flow routes multiple tokens through stages with arrival pulses & color
        s3=self.model["slides"][2]
        by3={o["name"]:o for o in s3["objects"]}
        st3={o["id"]:me.fresh_state() for o in s3["objects"]}
        cf=me.compile_choreography("cascade-flow",
                                   [by3["Packet"],by3["Firewall Blocker"],by3["Gateway"]],
                                   st3,by3,"cf",
                                   {"token_count":1,"hop_ms":400,"stage_color":"#06B6D4"})
        pulses=[e for e in cf if e["preset"]=="pulse"]
        colors=[e for e in cf if e["preset"]=="color"]
        self.assertEqual(len(pulses),2)
        self.assertEqual(len(colors),2)


if __name__=="__main__":
    unittest.main()
