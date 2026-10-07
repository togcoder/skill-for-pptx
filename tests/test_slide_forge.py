"""T025: Slide Forge spec -> deck -> motion -> QA."""
import json
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/"scripts"))

import forge_qa
import slide_forge as sf

EXAMPLES=ROOT/"examples"/"forge"


class SpecTests(unittest.TestCase):
    def test_schema_lists_every_layout(self):
        self.assertEqual(set(sf.schema()["slide"]["layout"]),set(sf.LAYOUTS))

    def test_errors_are_coded_with_fixes(self):
        spec={"theme":"neon","slides":[{"layout":"process","title":"x","steps":["a","b"]},
                                       {"layout":"chart","title":"c","chart":{"type":"radar","categories":["a"],
                                        "series":[{"name":"s","values":[1,2]}]}},
                                       {"layout":"wat"}]}
        codes={e["code"] for e in sf.validate_spec(spec)}
        self.assertEqual(codes,{"THEME_UNKNOWN","STRUCTURE","CHART_TYPE","CHART_DATA","LAYOUT_UNKNOWN"})
        self.assertTrue(all(e["fix"] for e in sf.validate_spec(spec)))

    def test_examples_validate(self):
        for path in EXAMPLES.glob("*.json"):
            self.assertEqual(sf.validate_spec(json.loads(path.read_text(encoding="utf-8"))),[],path.name)

    def test_theme_text_roles_meet_contrast(self):
        for name,t in sf.THEMES.items():
            for role in ("text","muted","accent","accent2"):
                for back in ("bg","surface"):
                    self.assertGreaterEqual(forge_qa.contrast(t[role],t[back]),4.5,f"{name} {role} on {back}")


class FitTests(unittest.TestCase):
    def test_long_word_never_fits_a_narrow_box(self):
        self.assertEqual(sf.text_lines("Customerization","Segoe UI Semibold",40,1.0,bold=True),float("inf"))

    def test_shrinks_until_text_fits(self):
        size,ok=sf.fit_size(["A fairly long sentence that has to wrap over a few lines"],"Segoe UI",40,4.0,1.2)
        self.assertTrue(ok)
        self.assertLess(size,40)


class BuildTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory()
        cls.out=Path(cls.tmp.name)/"coffee.pptx"
        spec=json.loads((EXAMPLES/"coffee_report.json").read_text(encoding="utf-8"))
        cls.verdict=sf.build(spec,cls.out)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_example_ships_clean(self):
        self.assertTrue(self.verdict["ok"],self.verdict["errors"])
        self.assertEqual([w for w in self.verdict["warnings"] if w["code"]!="TEXT_SHRUNK"],[])

    def test_every_slide_moves_and_no_picture_is_frozen(self):
        m=self.verdict["metrics"]
        self.assertEqual(m["motion_coverage"],1.0)
        self.assertEqual(m["pictures"],m["pictures_moving"])
        self.assertEqual(m["morph_transitions"],m["slides"]-1)

    def test_structures_are_choreographed(self):
        notes=self.verdict["director_notes"]
        self.assertEqual(notes["5"],"6 click(s), 10 motion beat(s)")  # notes walk 5 steps
        self.assertEqual(notes["7"],"7 click(s), 8 motion beat(s)")   # cycle: assemble, 4 spotlights, release
        self.assertIn("motion-graphic intro",notes["1"])

    def test_generated_art_is_labelled(self):
        self.assertEqual([g["object"] for g in self.verdict["generated_art"]],["Hero picture","Supporting picture"])


class QaTests(unittest.TestCase):
    def test_flags_overflow_collision_and_contrast(self):
        prs=Presentation()
        s=prs.slides.add_slide(prs.slide_layouts[6])
        a=s.shapes.add_textbox(Inches(1),Inches(1),Inches(2),Inches(0.4))
        a.text_frame.word_wrap=True
        a.text_frame.text="This sentence is far too long to fit inside such a small text box at this size"
        b=s.shapes.add_textbox(Inches(1.2),Inches(1.1),Inches(2),Inches(0.4))
        b.text_frame.text="Overlapping"
        from pptx.dml.color import RGBColor
        from pptx.util import Pt
        for box in (a,b):
            for r in box.text_frame.paragraphs[0].runs:
                r.font.size=Pt(18)
                r.font.color.rgb=RGBColor(0xEE,0xEE,0xEE)
        with tempfile.TemporaryDirectory() as td:
            path=Path(td)/"bad.pptx"
            prs.save(path)
            v=forge_qa.qa(path)
        codes={e["code"] for e in v["errors"]}
        self.assertTrue({"TEXT_OVERFLOW","TEXT_COLLISION","LOW_CONTRAST"}<=codes,codes)
        self.assertIn("MOTION_NONE",{w["code"] for w in v["warnings"]})


if __name__=="__main__":
    unittest.main()


class AssetTests(unittest.TestCase):
    """Offline: the cache is pre-seeded the way a real search leaves it."""

    def test_search_photo_icon_and_credits(self):
        import forge_assets as fa
        from PIL import Image
        from zipfile import ZipFile
        with tempfile.TemporaryDirectory() as td:
            cache=Path(td)/"cache"
            cache.mkdir()
            cand={"id":"ov-test","thumb":"x","url":"x","width":1600,"height":1000,"license":"by",
                  "license_url":"https://creativecommons.org/licenses/by/4.0/","attribution":"“Beans” by Ana, CC BY 4.0",
                  "landing":"https://example.org/beans","source":"wikimedia"}
            for orient in ("landscape","portrait"):
                (cache/f"search-openverse-{fa._key('beans',12,orient)}.json").write_text(json.dumps([cand]),encoding="utf-8")
            Image.new("RGB",(1600,1000),(120,80,40)).save(cache/"ov-test.jpg")
            for color in ("38BDF8","F472B6"):
                (cache/f"icon-lucide-star-{color}.svg").write_text(
                    '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24"><circle cx="12" cy="12" r="9" fill="#'+color+'"/></svg>')
            spec={"assets_cache":str(cache),"motion":{"morph":False},"slides":[
                {"layout":"title","title":"Beans","image":{"search":"beans"}},
                {"layout":"kpis","title":"Numbers","items":[{"value":"12%","label":"growth","icon":"lucide:star"}]},
                {"layout":"comparison","title":"A vs B","left":{"heading":"A","points":["x"],"icon":"lucide:star"},
                 "right":{"heading":"B","points":["y"],"icon":"lucide:star"}}]}
            out=Path(td)/"deck.pptx"
            v=sf.build(spec,out)
            self.assertTrue(v["ok"],v["errors"])
            self.assertEqual([c["attribution"] for c in v["image_credits"]],["“Beans” by Ana, CC BY 4.0"])
            self.assertEqual(len(v["icons"]),3)
            prs=Presentation(str(out))
            self.assertEqual(len(prs.slides),4)  # + credits slide
            self.assertIn("Beans",prs.slides[3].shapes[-1].text_frame.text)
            self.assertIn("Image: “Beans”",prs.slides[0].notes_slide.notes_text_frame.text)
            with ZipFile(out) as z:
                self.assertTrue(any(n.endswith(".svg") for n in z.namelist()))
                self.assertIn(b"svgBlip",z.read("ppt/slides/slide2.xml"))
                self.assertNotIn(b"p:timing",z.read("ppt/slides/slide4.xml"))  # [static] credits


class OfficeKitTests(unittest.TestCase):
    def test_every_office_kit_is_legible_in_both_modes(self):
        kits=json.loads(sf.KITS_PATH.read_text(encoding="utf-8"))["themes"]
        self.assertEqual(len(kits),11)
        for k in kits:
            for mode in ("light","dark"):
                t,fonts,_=sf.office_theme(k["name"],mode)
                for role in ("text","muted","accent","accent2"):
                    for back in ("bg","surface"):
                        self.assertGreaterEqual(forge_qa.contrast(t[role],t[back]),4.5,f"{k['name']} {mode} {role}/{back}")
                self.assertTrue(fonts["head"] and fonts["body"])

    def test_vietnamese_text_rejects_fonts_without_glyphs(self):
        spec={"theme":"office:Ion","slides":[{"layout":"section","title":"Chuyển đổi số"}]}
        self.assertIn("FONT_GLYPHS",{e["code"] for e in sf.validate_spec(spec)})
        spec["fonts"]={"head":"Segoe UI Semibold","body":"Segoe UI"}
        self.assertEqual(sf.validate_spec(spec),[])
