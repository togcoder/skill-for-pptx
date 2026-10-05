#!/usr/bin/env python3
"""Build the T023 compound-motion showcase fixture (synthetic data).

Development-only generator (python-pptx + Pillow). Committed output:
tests/fixtures/motion_showcase_deck.pptx. Each slide carries a structure that
calls for compound motion on existing objects: a PDCA cycle with a hub, a
hub-and-spoke ecosystem, a re-ranked priority list, a process with a token,
a 2x2 matrix, and two consecutive slides sharing a picture for Morph.
"""
import argparse
import io
import math
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

NAVY=RGBColor(0x1F,0x2A,0x44)
TEAL=RGBColor(0x0F,0x9D,0x8A)
SAND=RGBColor(0xF4,0xEF,0xE6)
CORAL=RGBColor(0xE0,0x6C,0x4F)
WHITE=RGBColor(0xFF,0xFF,0xFF)


def shape(slide,kind,x,y,w,h,text,fill,color=WHITE,size=18,name=None):
    s=slide.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb=fill
    s.line.fill.background()
    s.text_frame.text=text
    s.text_frame.word_wrap=True
    for p in s.text_frame.paragraphs:
        for r in p.runs:
            r.font.size=Pt(size)
            r.font.bold=True
            r.font.color.rgb=color
    if name:
        s.name=name
    return s


def picture_bytes():
    img=Image.new("RGB",(800,600),(244,239,230))
    d=ImageDraw.Draw(img)
    d.rectangle((80,260,720,520),fill=(31,42,68))
    d.polygon([(80,260),(400,90),(720,260)],fill=(15,157,138))
    d.rectangle((340,380,460,520),fill=(244,239,230))
    buf=io.BytesIO()
    img.save(buf,format="PNG")
    buf.seek(0)
    return buf


def build(path):
    prs=Presentation()
    prs.slide_width=Inches(13.333)
    prs.slide_height=Inches(7.5)
    title_layout,content,title_only=prs.slide_layouts[0],prs.slide_layouts[1],prs.slide_layouts[5]

    s=prs.slides.add_slide(title_layout)
    s.shapes.title.text="Operating Model 2027"
    s.placeholders[1].text="Northwind Roasters (synthetic data)"

    # PDCA cycle around a hub.
    s=prs.slides.add_slide(title_only)
    s.shapes.title.text="The improvement cycle"
    cx,cy,r=6.67,4.35,2.1
    shape(s,MSO_SHAPE.OVAL,cx-1.0,cy-1.0,2.0,2.0,"Customer value",NAVY,size=16,name="Hub")
    for i,label in enumerate(["Plan","Do","Check","Act"]):
        ang=-math.pi/2+i*math.pi/2
        x,y=cx+r*math.cos(ang),cy+r*math.sin(ang)
        shape(s,MSO_SHAPE.OVAL,x-0.8,y-0.8,1.6,1.6,label,TEAL,size=17,name=f"Stage {label}")
    s.notes_slide.notes_text_frame.text="Introduce the cycle, then walk Plan, Do, Check and Act."

    # Hub and six partners.
    s=prs.slides.add_slide(title_only)
    s.shapes.title.text="Partner ecosystem"
    shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,cx-1.4,cy-0.6,2.8,1.2,"Northwind platform",NAVY,size=18,name="Platform")
    for i,label in enumerate(["Growers","Importers","Cafés","Grocers","Couriers","Recyclers"]):
        ang=-math.pi/2+i*math.pi/3
        x,y=cx+2.6*math.cos(ang),cy+2.2*math.sin(ang)
        shape(s,MSO_SHAPE.OVAL,x-0.7,y-0.55,1.4,1.1,label,TEAL,size=14,name=f"Partner {label}")

    # Re-ranked priorities.
    s=prs.slides.add_slide(title_only)
    s.shapes.title.text="Priorities re-ranked for 2027"
    for i,label in enumerate(["Cost","Speed","Sustainability"]):
        shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,3.2,2.0+i*1.5,6.9,1.1,label,[NAVY,NAVY,CORAL][i],size=22,
              name=f"Priority {label}")

    # Process with a token.
    s=prs.slides.add_slide(title_only)
    s.shapes.title.text="Order journey"
    for i,label in enumerate(["Order","Roast","Pack","Deliver"]):
        shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,1.0+i*3.0,3.4,2.4,1.3,label,NAVY,size=20,name=f"Step {label}")
    shape(s,MSO_SHAPE.OVAL,0.35,3.75,0.6,0.6,"",CORAL,name="Order token")
    s.notes_slide.notes_text_frame.text="Follow one order through each step."

    # 2x2 matrix.
    s=prs.slides.add_slide(title_only)
    s.shapes.title.text="Market quadrants"
    labels=[("Grow","Premium cafés"),("Defend","Grocery retail"),("Explore","Subscriptions"),("Exit","Wholesale bulk")]
    for i,(head,body) in enumerate(labels):
        x=2.4+(i%2)*4.4
        y=1.9+(i//2)*2.6
        shape(s,MSO_SHAPE.ROUNDED_RECTANGLE,x,y,4.1,2.3,f"{head}\n{body}",[TEAL,NAVY,CORAL,NAVY][i],size=18,
              name=f"Quadrant {head}")

    # Morph pair: shared picture moves and shrinks.
    s=prs.slides.add_slide(title_only)
    s.shapes.title.text="Our roastery"
    pic=s.shapes.add_picture(picture_bytes(),Inches(0.8),Inches(1.8),Inches(7.0),Inches(5.25))
    pic.name="Roastery photo"
    s=prs.slides.add_slide(content)
    s.shapes.title.text="Our roastery"
    pic=s.shapes.add_picture(picture_bytes(),Inches(9.3),Inches(1.8),Inches(3.4),Inches(2.55))
    pic.name="Roastery photo"
    body=s.placeholders[1]
    body.left,body.top,body.width,body.height=Inches(0.8),Inches(1.8),Inches(8.0),Inches(4.8)
    body.text_frame.text="Capacity 1,200 t per year"
    body.text_frame.add_paragraph().text="Solar roof covers 40% of energy"
    body.text_frame.add_paragraph().text="Second site opens in 2027"

    path.parent.mkdir(parents=True,exist_ok=True)
    prs.save(path)


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument("output",type=Path,nargs="?",
                    default=Path(__file__).resolve().parents[1]/"tests"/"fixtures"/"motion_showcase_deck.pptx")
    args=ap.parse_args()
    build(args.output)
    print(args.output)


if __name__=="__main__":
    main()
