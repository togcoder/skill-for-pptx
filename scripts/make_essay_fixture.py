#!/usr/bin/env python3
"""Build the T022 synthetic text-box-only outline deck.

Development-only generator (needs python-pptx). Mirrors the *structure* of the
T009 water-report deck (no placeholders, title text boxes, labelled
alternatives, two-paragraph essay bodies, a wrapped multi-paragraph title) with
new synthetic Vietnamese content, so the regression does not depend on the
owner's source file.
"""
import argparse
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

LOREM_A=("Cây xanh đô thị không chỉ tạo bóng mát mà còn điều hòa vi khí hậu, giảm bụi mịn và "
         "giữ nước mưa. Mỗi hàng cây trưởng thành có thể hạ nhiệt độ mặt đường vài độ vào buổi trưa.")
LOREM_B=("Tuy nhiên, nhiều tuyến phố mới mở rộng lại chặt bỏ cây cũ mà không trồng bù tương xứng. "
         "Hệ quả là các đảo nhiệt xuất hiện dày hơn và chi phí làm mát của cư dân tăng lên rõ rệt.")


def _box(slide,x,y,w,h,paras,size,bold=False):
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=tb.text_frame
    tf.word_wrap=True
    for i,text in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=text
        for r in p.runs:
            r.font.size=Pt(size)
            r.font.bold=bold
    return tb


def build(path):
    prs=Presentation()
    prs.slide_width=Inches(13.333)
    prs.slide_height=Inches(7.5)
    blank=prs.slide_layouts[6]

    s=prs.slides.add_slide(blank)
    _box(s,1.2,1.0,10.9,0.5,["BÀI THUYẾT TRÌNH NGHỊ LUẬN"],27,True)
    _box(s,0.9,2.0,11.4,0.8,["CÂY XANH TRONG","ĐÔ THỊ HIỆN ĐẠI"],27,True)
    _box(s,1.2,3.4,10.9,0.8,["CHỦ ĐỀ: CON NGƯỜI VÀ","MÔI TRƯỜNG SỐNG"],22)

    s=prs.slides.add_slide(blank)
    _box(s,0.2,0.15,12.9,0.4,["I. MỞ BÀI"],20,True)
    _box(s,0.5,0.8,12.0,0.3,["Cách 1: Đi từ trải nghiệm đời thường"],16,True)
    _box(s,0.8,1.2,11.8,2.2,[LOREM_A],16.8)
    _box(s,0.5,3.75,12.0,0.3,["Cách 2: Đi từ số liệu"],16,True)
    _box(s,0.8,4.05,11.8,2.4,[LOREM_B],16.8)

    s=prs.slides.add_slide(blank)
    _box(s,0.2,0.15,12.9,0.4,["II. THÂN BÀI"],20,True)
    _box(s,0.3,0.5,12.8,0.3,["1. Vai trò của cây xanh"],13,True)
    _box(s,0.7,0.95,12.1,5.4,[LOREM_A,"",LOREM_B],17)

    s=prs.slides.add_slide(blank)
    _box(s,0.2,0.15,12.9,0.4,["2. Ý kiến trái chiều và phản bác"],20,True)
    _box(s,0.7,0.95,12.1,5.6,[LOREM_B,"",LOREM_A],17)

    s=prs.slides.add_slide(blank)
    _box(s,0.2,0.15,12.9,0.4,["III. KẾT BÀI"],20,True)
    _box(s,0.3,0.5,12.8,0.3,["CÁC CÁCH KẾT BÀI"],13,True)
    _box(s,0.5,1.05,12.0,0.3,["Cách 1: Khẳng định"],16,True)
    _box(s,0.8,1.35,11.8,2.2,[LOREM_A],16.5)
    _box(s,0.5,3.8,12.0,0.3,["Cách 2: Gợi mở"],16,True)
    _box(s,0.8,4.2,11.8,2.2,[LOREM_B],16.5)

    path.parent.mkdir(parents=True,exist_ok=True)
    prs.save(path)


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument("output",type=Path,nargs="?",
                    default=Path(__file__).resolve().parents[1]/"tests"/"fixtures"/"essay_outline_deck.pptx")
    args=ap.parse_args()
    build(args.output)
    print(args.output)


if __name__=="__main__":
    main()
