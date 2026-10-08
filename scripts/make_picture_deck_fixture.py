#!/usr/bin/env python3
"""Build the T024 picture/motion-graphics fixture (synthetic data).

Development-only generator (python-pptx + Pillow). Committed output:
tests/fixtures/picture_deck.pptx. It mimics what a prompt-to-deck generator
typically produces: text boxes, pictures and thin accent bars, no
placeholders, no animation.

  1. title over a full-bleed backdrop picture, with an accent bar
  2. heading + picture + two-paragraph text, with an accent underline
  3. a single picture as the slide's content
"""
import argparse
import io
from pathlib import Path

from PIL import Image, ImageDraw

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

NAVY=RGBColor(0x1F,0x2A,0x44)
CORAL=RGBColor(0xE0,0x6C,0x4F)


def photo(hue):
    """Distinct synthetic picture per slide (identical bytes would read as a Morph pair)."""
    img=Image.new("RGB",(800,600),hue)
    d=ImageDraw.Draw(img)
    d.ellipse((220,120,580,480),fill=(244,239,230))
    d.rectangle((0,470,800,600),fill=(31,42,68))
    buf=io.BytesIO()
    img.save(buf,format="PNG")
    buf.seek(0)
    return buf


def text(slide,x,y,w,h,paragraphs,size,name,bold=False):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    box.name=name
    tf=box.text_frame
    tf.word_wrap=True
    for i,t in enumerate(paragraphs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=t
        for r in p.runs:
            r.font.size=Pt(size)
            r.font.bold=bold
            r.font.color.rgb=NAVY
    return box


def bar(slide,x,y,w,h,name):
    s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb=CORAL
    s.line.fill.background()
    s.name=name
    return s


def build(path):
    prs=Presentation()
    prs.slide_width=Inches(13.333)
    prs.slide_height=Inches(7.5)
    blank=prs.slide_layouts[6]

    s=prs.slides.add_slide(blank)
    pic=s.shapes.add_picture(photo((15,157,138)),0,0,prs.slide_width,prs.slide_height)
    pic.name="Backdrop"
    text(s,0.8,2.6,9.0,1.4,["Coffee that travels well"],44,"Title",bold=True)
    bar(s,0.8,4.1,3.2,0.08,"Accent bar")

    s=prs.slides.add_slide(blank)
    text(s,0.8,0.5,11.0,0.9,["Why our roastery"],32,"Heading",bold=True)
    bar(s,0.8,1.4,2.0,0.06,"Heading underline")
    pic=s.shapes.add_picture(photo((224,108,79)),Inches(7.3),Inches(1.9),Inches(5.2),Inches(3.9))
    pic.name="Roastery photo"
    text(s,0.8,1.9,6.0,3.9,["Roasted within 48 hours of every order.",
                            "Solar roof covers 40% of the energy we use."],20,"Body")

    s=prs.slides.add_slide(blank)
    pic=s.shapes.add_picture(photo((120,90,160)),Inches(2.2),Inches(0.9),Inches(8.9),Inches(5.7))
    pic.name="Harvest photo"

    path.parent.mkdir(parents=True,exist_ok=True)
    prs.save(path)


def main():
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument("output",type=Path,nargs="?",
                    default=Path(__file__).resolve().parents[1]/"tests"/"fixtures"/"picture_deck.pptx")
    args=ap.parse_args()
    build(args.output)
    print(args.output)


if __name__=="__main__":
    main()
