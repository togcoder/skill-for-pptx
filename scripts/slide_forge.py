#!/usr/bin/env python3
"""Slide Forge — an agent-facing deck compiler (T025).

An AI writes a semantic deck spec (JSON); Forge turns it into a finished,
editable, motion-directed PPTX and a machine-readable QA verdict. Nothing here
is meant for a human to click through: every input is declarative, every
failure is a coded error with a fix hint, and the loop is
``spec -> build -> qa -> (repair spec) -> build``.

    slide_forge.py schema                      # spec contract as JSON
    slide_forge.py build SPEC.json -o OUT.pptx # layout + art + motion + QA
    slide_forge.py qa OUT.pptx                 # QA only (any PPTX)
    slide_forge.py render OUT.pptx -d DIR      # PowerPoint render -> MP4 + per-slide strips

Layouts: title, section, statement, bullets, image, kpis, process, cycle,
chart, comparison, quote, closing. Layout geometry is chosen so the Motion
Director recognises the structure (cards, process rows, cycles, charts) and
choreographs it; pictures get Ken Burns and accents draw in (T024). Shared
``!!`` background orbs move between slides so Morph adds a motion-graphic
layer across the deck.
"""
import argparse
import copy
import hashlib
import io
import json
import math
import random
import sys
from pathlib import Path

from lxml import etree as E
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

sys.path.insert(0,str(Path(__file__).resolve().parent))

VERSION="forge-0.1"
W,H=13.333,7.5
MARGIN=0.7
THEMES={
    "midnight":{"bg":"0F172A","surface":"1E293B","text":"F8FAFC","muted":"94A3B8","accent":"38BDF8","accent2":"F472B6"},
    "paper":{"bg":"FAF7F2","surface":"FFFFFF","text":"1F2A44","muted":"5B6472","accent":"B4532F","accent2":"0F766E"},
    "forest":{"bg":"0B1F17","surface":"13342A","text":"ECFDF5","muted":"A7B5AE","accent":"34D399","accent2":"FBBF24"},
    "sunrise":{"bg":"FFF7ED","surface":"FFFFFF","text":"431407","muted":"7C5A48","accent":"C2410C","accent2":"7C3AED"},
}
FONTS={"head":"Segoe UI Semibold","body":"Segoe UI"}
LAYOUTS={
    "title":{"required":["title"],"optional":["subtitle","image","kicker"]},
    "section":{"required":["title"],"optional":["kicker"]},
    "statement":{"required":["text"],"optional":["kicker"]},
    "bullets":{"required":["title","bullets"],"optional":["image"]},
    "image":{"required":["image"],"optional":["title","caption"]},
    "kpis":{"required":["title","items"],"optional":["takeaway"]},
    "process":{"required":["title","steps"],"optional":["takeaway"]},
    "cycle":{"required":["title","items"],"optional":["center"]},
    "chart":{"required":["title","chart"],"optional":["takeaway"]},
    "comparison":{"required":["title","left","right"],"optional":[]},
    "quote":{"required":["quote"],"optional":["author"]},
    "closing":{"required":["title"],"optional":["subtitle"]},
}
CHART_TYPES={"column":XL_CHART_TYPE.COLUMN_CLUSTERED,"bar":XL_CHART_TYPE.BAR_CLUSTERED,
             "line":XL_CHART_TYPE.LINE_MARKERS,"pie":XL_CHART_TYPE.PIE,"doughnut":XL_CHART_TYPE.DOUGHNUT,
             "stacked":XL_CHART_TYPE.COLUMN_STACKED,"area":XL_CHART_TYPE.AREA}
LIMITS={"bullets":6,"kpis":4,"process":6,"cycle":6,"points":5}


# ---------------------------------------------------------------------------
# Spec contract


def schema():
    return {
        "version":VERSION,
        "deck":{"title":"str (optional)","assets_cache":"dir for downloaded photos/icons (default ~/.cache/slide-forge)",
                "credits_slide":"bool, closing slide attributing licensed photos (default true)","theme":f"one of {sorted(THEMES)} or an object with the same color keys (hex)",
                "fonts":{"head":"font family","body":"font family"},
                "motion":{"style":"subtle|modern|bold|cinematic (default modern)","counters":"bool, count up KPI values (default true)",
                          "morph":"bool, background orbs glide between slides (default true)"},
                "slides":"list of slide objects"},
        "slide":{"layout":{k:v for k,v in LAYOUTS.items()},"notes":"speaker notes; sequence words (first, then, finally) make process steps one click each"},
        "fields":{"image":"path, or {\"search\":\"english keywords\", \"pick\":n?, \"orientation\"?} for a licensed online photo "
                          "(Openverse CC0/BY/BY-SA, or Pexels with PEXELS_API_KEY; preview with `forge_assets.py photos`), "
                          "or {\"generate\":\"seed words\"} for on-theme abstract art (labelled generated)",
                  "icon":"Iconify name like \"lucide:coffee\" or plain keywords; on kpis items and comparison left/right",
                  "bullets":f"list[str] <= {LIMITS['bullets']}",
                  "items (kpis)":f"list[{{value:str, label:str}}] <= {LIMITS['kpis']}",
                  "steps":f"list[str] 3..{LIMITS['process']}",
                  "items (cycle)":f"list[str] 3..{LIMITS['cycle']}",
                  "chart":f"{{type:{sorted(CHART_TYPES)}, categories:list[str], series:list[{{name, values:list[number]}}], number_format?}}",
                  "left/right":f"{{heading:str, points:list[str] <= {LIMITS['points']}}}"},
        "contract":["Text is never truncated: Forge shrinks type to fit and reports TEXT_SHRUNK; below the floor it fails TEXT_OVERFLOW.",
                    "Every error has code, slide, object, detail and fix; repair the spec, not the PPTX.",
                    "Generated art is abstract and labelled; it never stands in for factual imagery."],
    }


def validate_spec(spec):
    errs=[]
    def err(code,slide,detail,fix):
        errs.append({"code":code,"slide":slide,"detail":detail,"fix":fix})
    if not isinstance(spec,dict) or not isinstance(spec.get("slides"),list) or not spec["slides"]:
        err("SPEC_SHAPE",None,"spec needs a nonempty slides list","see `slide_forge.py schema`")
        return errs
    theme=spec.get("theme","midnight")
    if isinstance(theme,str) and theme not in THEMES:
        err("THEME_UNKNOWN",None,f"theme {theme!r}",f"use one of {sorted(THEMES)}")
    style=(spec.get("motion") or {}).get("style","modern")
    if style not in ("subtle","modern","bold","cinematic"):
        err("MOTION_STYLE",None,f"style {style!r}","use subtle|modern|bold|cinematic")
    for i,s in enumerate(spec["slides"],1):
        lay=s.get("layout") if isinstance(s,dict) else None
        if lay not in LAYOUTS:
            err("LAYOUT_UNKNOWN",i,f"layout {lay!r}",f"use one of {sorted(LAYOUTS)}")
            continue
        for key in LAYOUTS[lay]["required"]:
            if not s.get(key):
                err("FIELD_MISSING",i,f"{lay} needs {key!r}","add the field")
        n=lambda k:len(s.get(k) or [])
        if lay=="bullets" and n("bullets")>LIMITS["bullets"]:
            err("TOO_DENSE",i,f"{n('bullets')} bullets",f"split into two slides or keep <= {LIMITS['bullets']}")
        if lay=="kpis" and not 1<=n("items")<=LIMITS["kpis"]:
            err("TOO_DENSE",i,f"{n('items')} KPIs",f"use 1..{LIMITS['kpis']}")
        if lay=="process" and not 3<=n("steps")<=LIMITS["process"]:
            err("STRUCTURE",i,f"{n('steps')} steps",f"a process needs 3..{LIMITS['process']} steps")
        if lay=="cycle" and not 3<=n("items")<=LIMITS["cycle"]:
            err("STRUCTURE",i,f"{n('items')} cycle items",f"a cycle needs 3..{LIMITS['cycle']} items")
        if lay=="chart" and s.get("chart"):
            c=s["chart"]
            if c.get("type") not in CHART_TYPES:
                err("CHART_TYPE",i,f"chart type {c.get('type')!r}",f"use one of {sorted(CHART_TYPES)}")
            cats=c.get("categories") or []
            for ser in c.get("series") or []:
                if len(ser.get("values") or [])!=len(cats):
                    err("CHART_DATA",i,f"series {ser.get('name')!r} has {len(ser.get('values') or [])} values for {len(cats)} categories",
                        "give one value per category")
            if not c.get("series"):
                err("CHART_DATA",i,"chart has no series","add series:[{name, values}]")
        if lay=="comparison":
            for side in ("left","right"):
                if s.get(side) and len(s[side].get("points") or [])>LIMITS["points"]:
                    err("TOO_DENSE",i,f"{side} has too many points",f"keep <= {LIMITS['points']}")
        img=s.get("image")
        if isinstance(img,str) and not Path(img).is_file():
            err("IMAGE_MISSING",i,f"no file {img!r}","fix the path, or use {\"search\":\"...\"} / {\"generate\":\"...\"}")
        if isinstance(img,dict) and not (img.get("search") or img.get("generate")):
            err("IMAGE_REF",i,"image object needs search or generate","{\"search\": \"words\", \"pick\": n?}")
    return errs


# ---------------------------------------------------------------------------
# Text fitting


_FONT_FILES={"segoe ui":["segoeui.ttf","DejaVuSans.ttf"],"segoe ui bold":["segoeuib.ttf","DejaVuSans-Bold.ttf"],
             "segoe ui semibold":["seguisb.ttf","segoeuib.ttf","DejaVuSans-Bold.ttf"],
             "segoe ui semibold bold":["segoeuib.ttf","DejaVuSans-Bold.ttf"],
             "arial":["arial.ttf","LiberationSans-Regular.ttf"],"arial bold":["arialbd.ttf","LiberationSans-Bold.ttf"],
             "calibri":["calibri.ttf","Carlito-Regular.ttf"],"calibri bold":["calibrib.ttf","Carlito-Bold.ttf"]}
SAFETY=1.06  # renderer kerning/hinting differs slightly from Pillow's metrics
_FONT_DIRS=[Path("C:/Windows/Fonts"),Path("/usr/share/fonts"),Path("/Library/Fonts"),Path.home()/".fonts"]
_font_cache={}
_path_cache={}


def _font_path(family):
    fam=family.lower()
    if fam not in _path_cache:
        _path_cache[fam]=None
        for name in _FONT_FILES.get(fam,[]):
            for d in _FONT_DIRS:
                if (d/name).is_file():
                    _path_cache[fam]=d/name
                elif d.is_dir() and d.name!="Fonts":
                    _path_cache[fam]=next(d.rglob(name),None)
                if _path_cache[fam]:
                    return _path_cache[fam]
    return _path_cache[fam]


def _font(family,size_pt):
    key=(family.lower(),size_pt)
    if key not in _font_cache:
        path=_font_path(family)
        _font_cache[key]=ImageFont.truetype(str(path),size=max(1,round(size_pt*4))) if path else None
    return _font_cache[key]


def text_lines(text,family,size_pt,width_in,bold=False):
    """Lines after word wrap, or inf when a single word is wider than the box
    (PowerPoint would break it mid-word). Measured with the real font file
    when present; otherwise 0.55 em per character (ponytail: approximate)."""
    fam=family+" bold" if bold and (family.lower()+" bold") in _FONT_FILES else family
    font=_font(fam,size_pt)
    width_px=width_in*72*4  # font loaded at 4 px per pt
    def w(s):
        return (font.getlength(s) if font else len(s)*size_pt*4*(0.6 if bold else 0.55))*SAFETY
    if any(w(word)>width_px for word in str(text).split()):
        return math.inf
    lines=0
    for para in str(text).split("\n"):
        words=para.split() or [""]
        cur=""
        n=1
        for word in words:
            trial=(cur+" "+word).strip()
            if w(trial)<=width_px or not cur:
                cur=trial
            else:
                n+=1
                cur=word
        lines+=n
    return lines


def fit_size(paragraphs,family,size_pt,box_w,box_h,floor=12,spacing=1.18,gap_pt=6,bold=False):
    """Largest size <= size_pt whose wrapped text fits the box (inches)."""
    size=size_pt
    while True:
        lines=sum(text_lines(p,family,size,box_w-0.2,bold) for p in paragraphs)
        height=(lines*size*spacing+gap_pt*(len(paragraphs)-1))/72+0.15
        if height<=box_h or size<=floor:
            return size,height<=box_h
        size-=1


# ---------------------------------------------------------------------------
# Generated art


def _hex(c):
    return tuple(int(c[i:i+2],16) for i in (0,2,4))


def generated_art(seed,theme,w=1600,h=1000):
    """On-theme abstract art: soft gradient field with blurred orbs and arcs."""
    rnd=random.Random(hashlib.sha256(seed.encode()).hexdigest())
    a,b,bg=_hex(theme["accent"]),_hex(theme["accent2"]),_hex(theme["surface"])
    img=Image.new("RGB",(w,h),bg)
    px=Image.linear_gradient("L").resize((w,h)).rotate(rnd.choice([0,45,90,135]),expand=False)
    img=Image.composite(Image.new("RGB",(w,h),tuple((x*3+y)//4 for x,y in zip(bg,a))),img,px)
    layer=Image.new("RGBA",(w,h),(0,0,0,0))
    d=ImageDraw.Draw(layer)
    for _ in range(7):
        r=rnd.randint(h//6,h//2)
        x,y=rnd.randint(0,w),rnd.randint(0,h)
        col=rnd.choice([a,b])
        d.ellipse((x-r,y-r,x+r,y+r),fill=col+(rnd.randint(70,150),))
    layer=layer.filter(ImageFilter.GaussianBlur(h//14))
    img=Image.alpha_composite(img.convert("RGBA"),layer)
    d=ImageDraw.Draw(img)
    for i in range(4):
        r=h//3+i*h//10
        cx,cy=rnd.randint(w//4,3*w//4),rnd.randint(h//4,3*h//4)
        d.arc((cx-r,cy-r,cx+r,cy+r),rnd.randint(0,180),rnd.randint(200,340),fill=(255,255,255,90),width=3)
    buf=io.BytesIO()
    img.convert("RGB").save(buf,format="PNG")
    buf.seek(0)
    return buf


# ---------------------------------------------------------------------------
# Builder


class Builder:
    def __init__(self,spec):
        t=spec.get("theme","midnight")
        self.theme=dict(THEMES[t]) if isinstance(t,str) else {**THEMES["midnight"],**t}
        self.fonts={**FONTS,**(spec.get("fonts") or {})}
        self.motion={"style":"modern","counters":True,"morph":True,**(spec.get("motion") or {})}
        self.prs=Presentation()
        self.prs.slide_width=Inches(W)
        self.prs.slide_height=Inches(H)
        self.blank=self.prs.slide_layouts[6]
        self.receipts=[]  # TEXT_SHRUNK etc.
        self.generated=[]
        self.credits=[]
        self.icons=[]
        self.cache=Path(spec.get("assets_cache") or Path.home()/".cache"/"slide-forge")

    def rgb(self,key):
        return RGBColor.from_string(self.theme.get(key,key))

    # -- primitives ---------------------------------------------------------

    def text(self,slide,x,y,w,h,paragraphs,size,name,role="body",color="text",bold=False,align=None,
             anchor=MSO_ANCHOR.TOP,floor=12,idx=None,gap=6):
        if isinstance(paragraphs,str):
            paragraphs=[paragraphs]
        family=self.fonts["head" if role=="head" else "body"]
        fitted,ok=fit_size(paragraphs,family,size,w,h,floor,gap_pt=gap,bold=bold)
        if fitted<size:
            self.receipts.append({"code":"TEXT_SHRUNK","slide":idx,"object":name,"detail":f"{size}->{fitted} pt",
                                  "fix":"shorten the text if the smaller size hurts hierarchy"})
        if not ok:
            self.receipts.append({"code":"TEXT_OVERFLOW","slide":idx,"object":name,"detail":f"does not fit at {floor} pt",
                                  "fix":"shorten the text or split the slide","severity":"error"})
        box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        box.name=name
        tf=box.text_frame
        tf.word_wrap=True
        tf.vertical_anchor=anchor
        for m in ("margin_left","margin_right"):
            setattr(tf,m,Inches(0.1))
        for i,t in enumerate(paragraphs):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.text=str(t)
            if align is not None:
                p.alignment=align
            if i:
                p.space_before=Pt(gap)
            for r in p.runs:
                r.font.size=Pt(fitted)
                r.font.bold=bold
                r.font.name=family
                r.font.color.rgb=self.rgb(color)
        return box

    def rect(self,slide,x,y,w,h,fill,name,shape=MSO_SHAPE.RECTANGLE,line=None):
        s=slide.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
        s.name=name
        s.fill.solid()
        s.fill.fore_color.rgb=self.rgb(fill)
        if line:
            s.line.color.rgb=self.rgb(line)
            s.line.width=Pt(1)
        else:
            s.line.fill.background()
        s.shadow.inherit=False
        if shape==MSO_SHAPE.ROUNDED_RECTANGLE:
            s.adjustments[0]=0.08
        return s

    def picture(self,slide,img,x,y,w,h,name,idx):
        """Place a picture cropped to fill the box (no distortion)."""
        if isinstance(img,dict) and "search" in img:
            import forge_assets
            src,cand=forge_assets.resolve_photo(img,self.cache,w/h)
            src=str(src)
            self.credits.append({"slide":idx,"object":name,**{k:cand[k] for k in
                                 ("id","attribution","license","license_url","landing","source")}})
        elif isinstance(img,dict):
            src=generated_art(img.get("generate","")+str(idx),self.theme)
            self.generated.append({"slide":idx,"object":name,"seed":img.get("generate","")})
        else:
            src=img
        pic=slide.shapes.add_picture(src,Inches(x),Inches(y),Inches(w),Inches(h))
        pic.name=name
        iw,ih=pic.image.size
        box,ratio=w/h,iw/ih
        if ratio>box:
            c=(1-box/ratio)/2
            pic.crop_left=pic.crop_right=c
        else:
            c=(1-ratio/box)/2
            pic.crop_top=pic.crop_bottom=c
        return pic

    def icon(self,slide,ref,x,y,size,name,color="accent"):
        """Vector icon (Iconify, open-source sets) as an SVG picture: PowerPoint
        2016+ draws the SVG; older readers get a transparent PNG fallback."""
        import forge_assets
        from pptx.opc.constants import RELATIONSHIP_TYPE as RT
        from pptx.opc.package import Part
        svg_path,resolved=forge_assets.fetch_icon(ref,self.theme.get(color,color),self.cache)
        buf=io.BytesIO()
        Image.new("RGBA",(256,256),(0,0,0,0)).save(buf,format="PNG")
        buf.seek(0)
        pic=slide.shapes.add_picture(buf,Inches(x),Inches(y),Inches(size),Inches(size))
        pic.name=name
        pkg=slide.part.package
        part=Part(pkg.next_partname("/ppt/media/image%d.svg"),"image/svg+xml",pkg,svg_path.read_bytes())
        rid=slide.part.relate_to(part,RT.IMAGE)
        blip=pic._element.find(".//"+qn("a:blip"))
        ext=E.SubElement(E.SubElement(blip,qn("a:extLst")),qn("a:ext"),uri="{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
        E.SubElement(ext,"{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip",
                     {qn("r:embed"):rid},nsmap={"asvg":"http://schemas.microsoft.com/office/drawing/2016/SVG/main"})
        self.icons.append({"object":name,"icon":resolved})
        return pic

    def background(self,slide,idx,n):
        fill=slide.background.fill
        fill.solid()
        fill.fore_color.rgb=self.rgb("bg")
        if not self.motion["morph"]:
            return
        # Two soft orbs with fixed names sit on opposite corners and rotate a
        # corner per slide, so Morph glides them around the frame, never over content.
        corners=[(0,0),(W,0),(W,H),(0,H)]
        for k,(key,r) in enumerate((("accent",2.0),("accent2",1.4))):
            cx,cy=corners[(idx+2*k)%4]
            orb=self.rect(slide,cx-r,cy-r,2*r,2*r,key,f"!!orb-{k+1}",MSO_SHAPE.OVAL)
            _alpha(orb,12)

    def heading(self,slide,title,idx,kicker=None):
        y=0.45
        if kicker:
            self.text(slide,MARGIN,0.32,W-2*MARGIN,0.4,[kicker.upper()],13,"Kicker",color="accent",bold=True,idx=idx)
            y=0.62
        self.text(slide,MARGIN,y,W-2*MARGIN,0.85,[title],36,"Title",role="head",bold=True,idx=idx)
        self.rect(slide,MARGIN+0.1,y+0.92,1.4,0.07,"accent","Title accent")
        return y+1.25

    # -- layouts ------------------------------------------------------------

    def l_title(self,slide,s,idx):
        img=s.get("image") or {"generate":s["title"]}
        self.picture(slide,img,W/2,0,W/2,H,"Hero picture",idx)
        if s.get("kicker"):
            self.text(slide,MARGIN,2.0,W/2-1.2,0.45,[s["kicker"].upper()],14,"Kicker",color="accent",bold=True,idx=idx)
        self.text(slide,MARGIN,2.45,W/2-1.2,2.2,[s["title"]],46,"Title",role="head",bold=True,
                  anchor=MSO_ANCHOR.BOTTOM,idx=idx)
        self.rect(slide,MARGIN+0.1,4.8,1.8,0.08,"accent","Title accent")
        if s.get("subtitle"):
            self.text(slide,MARGIN,5.05,W/2-1.2,1.2,[s["subtitle"]],20,"Subtitle",color="muted",idx=idx)

    def l_section(self,slide,s,idx):
        if s.get("kicker"):
            self.text(slide,MARGIN,2.5,W-2*MARGIN,0.5,[s["kicker"].upper()],16,"Kicker",color="accent",bold=True,idx=idx)
        self.text(slide,MARGIN,3.0,W-2*MARGIN,1.4,[s["title"]],44,"Title",role="head",bold=True,idx=idx)
        self.rect(slide,MARGIN+0.1,4.5,2.4,0.08,"accent","Section accent")

    def l_statement(self,slide,s,idx):
        if s.get("kicker"):
            self.text(slide,1.4,1.2,W-2.8,0.5,[s["kicker"].upper()],15,"Kicker",color="accent",bold=True,idx=idx)
        self.rect(slide,1.2,1.85,0.08,3.6,"accent","Statement bar")
        self.text(slide,1.5,1.7,W-3.0,3.9,[s["text"]],54,"Statement",role="head",bold=True,
                  anchor=MSO_ANCHOR.MIDDLE,idx=idx)

    def l_bullets(self,slide,s,idx):
        top=self.heading(slide,s["title"],idx,s.get("kicker"))
        width=W-2*MARGIN
        if s.get("image"):
            width=W*0.52-MARGIN
            self.picture(slide,s["image"],W*0.56,top,W*0.44-MARGIN,H-top-0.7,"Supporting picture",idx)
        self.text(slide,MARGIN,top+0.1,width,H-top-0.8,["• "+b for b in s["bullets"]],28,"Points",
                  anchor=MSO_ANCHOR.MIDDLE,gap=18,idx=idx)

    def l_image(self,slide,s,idx):
        if s.get("title"):
            top=self.heading(slide,s["title"],idx)
        else:
            top=0.6
        cap=0.7 if s.get("caption") else 0
        self.picture(slide,s["image"],1.4,top,W-2.8,H-top-0.5-cap,"Feature picture",idx)
        if cap:
            self.text(slide,1.4,H-0.5-cap+0.1,W-2.8,cap-0.1,[s["caption"]],15,"Caption",color="muted",
                      align=PP_ALIGN.CENTER,idx=idx)

    def l_kpis(self,slide,s,idx):
        top=self.heading(slide,s["title"],idx)
        items=s["items"]
        n=len(items)
        gap=0.35
        cw=(W-2*MARGIN-gap*(n-1))/n
        ch=3.0+(0.55 if any(t.get("icon") for t in items) else 0)
        y=top+(H-top-0.5-ch-(1.1 if s.get("takeaway") else 0))/2
        # One size for every value so the row reads as a set.
        vsize=min(fit_size([it["value"]],self.fonts["head"],64,cw-0.4,1.2,floor=24,bold=True)[0] for it in items)
        for i,it in enumerate(items):
            x=MARGIN+i*(cw+gap)
            self.rect(slide,x,y,cw,ch,"surface",f"KPI card {i+1}",MSO_SHAPE.ROUNDED_RECTANGLE)
            dy=0
            if any(t.get("icon") for t in items):
                dy=0.35
                if it.get("icon"):
                    self.icon(slide,it["icon"],x+cw/2-0.3,y+0.22,0.6,f"KPI icon {i+1}")
            self.text(slide,x+0.2,y+0.3+dy*1.6,cw-0.4,1.2,[it["value"]],vsize,f"KPI value {i+1}",role="head",color="accent",
                      bold=True,align=PP_ALIGN.CENTER,idx=idx)
            self.text(slide,x+0.2,y+1.75+dy*1.6,cw-0.4,1.0-dy*0.6,[it["label"]],20,f"KPI label {i+1}",color="muted",
                      align=PP_ALIGN.CENTER,idx=idx)
        if s.get("takeaway"):
            self.text(slide,MARGIN,y+ch+0.35,W-2*MARGIN,0.8,[s["takeaway"]],24,"Takeaway",bold=True,
                      align=PP_ALIGN.CENTER,idx=idx)

    def l_process(self,slide,s,idx):
        top=self.heading(slide,s["title"],idx)
        steps=s["steps"]
        n=len(steps)
        gap=0.55
        bw=(W-2*MARGIN-gap*(n-1))/n
        bh=2.2
        y=top+(H-top-0.5-bh-(1.2 if s.get("takeaway") else 0))/2
        for i,st in enumerate(steps):
            x=MARGIN+i*(bw+gap)
            box=self.rect(slide,x,y,bw,bh,"surface" if i<n-1 else "accent",f"Step {i+1}",MSO_SHAPE.ROUNDED_RECTANGLE)
            tf=box.text_frame
            tf.word_wrap=True
            fam=self.fonts["head"]
            size,_=fit_size([st],fam,24,bw-0.2,bh-0.2,bold=True)
            tf.text=st
            for p in tf.paragraphs:
                p.alignment=PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.size=Pt(size)
                    r.font.bold=True
                    r.font.name=fam
                    r.font.color.rgb=self.rgb("text" if i<n-1 else "bg")
            if i:
                c=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x-gap+0.08),Inches(y+bh/2),
                                             Inches(x-0.08),Inches(y+bh/2))
                c.name=f"Link {i}"
                c.line.color.rgb=self.rgb("accent")
                c.line.width=Pt(2.5)
                _arrow(c)
        if s.get("takeaway"):
            self.text(slide,MARGIN,y+bh+0.5,W-2*MARGIN,0.8,[s["takeaway"]],24,"Takeaway",bold=True,
                      align=PP_ALIGN.CENTER,idx=idx)

    def l_cycle(self,slide,s,idx):
        top=self.heading(slide,s["title"],idx)
        items=s["items"]
        n=len(items)
        cx,cy=W/2,(top+H-0.3)/2
        r=min(2.3,(H-top-0.4)/2-0.85)
        d=1.65
        if s.get("center"):
            hub=self.rect(slide,cx-0.85,cy-0.85,1.7,1.7,"accent","Hub",MSO_SHAPE.OVAL)
            _shape_text(self,hub,s["center"],20,"bg")
        for i,it in enumerate(items):
            ang=-math.pi/2+2*math.pi*i/n
            x,y=cx+r*1.2*math.cos(ang),cy+r*math.sin(ang)  # near-circular: the Director detects cycles by radius spread
            node=self.rect(slide,x-d/2,y-d/2,d,d,"surface","Stage "+str(i+1),MSO_SHAPE.OVAL,line="accent")
            _shape_text(self,node,it,18,"text")

    def l_chart(self,slide,s,idx):
        top=self.heading(slide,s["title"],idx)
        c=s["chart"]
        data=CategoryChartData(number_format=c.get("number_format","General"))
        data.categories=c["categories"]
        for ser in c["series"]:
            data.add_series(ser["name"],ser["values"])
        cw=W-2*MARGIN if not s.get("takeaway") else W*0.62-MARGIN
        gf=slide.shapes.add_chart(CHART_TYPES[c["type"]],Inches(MARGIN),Inches(top+0.1),Inches(cw),Inches(H-top-0.6),data)
        gf.name="Chart"
        ch=gf.chart
        ch.font.size=Pt(18)
        ch.font.name=self.fonts["body"]
        ch.font.color.rgb=self.rgb("muted")
        ch.has_legend=len(c["series"])>1 or c["type"] in ("pie","doughnut")
        if ch.has_legend:
            ch.legend.position=XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout=False
        palette=[self.theme["accent"],self.theme["accent2"],self.theme["muted"],self.theme["text"]]
        plot=ch.plots[0]
        if c["type"] in ("pie","doughnut"):
            for j,pt in enumerate(plot.series[0].points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb=RGBColor.from_string(palette[j%len(palette)])
        else:
            for j,ser in enumerate(plot.series):
                col=RGBColor.from_string(palette[j%len(palette)])
                if c["type"]=="line":
                    ser.format.line.color.rgb=col
                    ser.format.line.width=Pt(3)
                else:
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb=col
            for ax in (ch.category_axis,ch.value_axis):
                ax.format.line.color.rgb=self.rgb("muted")
                ax.tick_labels.font.color.rgb=self.rgb("muted")
            ch.value_axis.has_major_gridlines=False
        if s.get("takeaway"):
            x=W*0.62+0.3
            self.rect(slide,x,top+0.6,0.07,2.6,"accent","Takeaway bar")
            self.text(slide,x+0.25,top+0.5,W-x-0.25-MARGIN,2.9,[s["takeaway"]],30,"Takeaway",role="head",
                      bold=True,anchor=MSO_ANCHOR.MIDDLE,idx=idx)

    def l_comparison(self,slide,s,idx):
        top=self.heading(slide,s["title"],idx)
        gap=0.5
        cw=(W-2*MARGIN-gap)/2
        chh=H-top-0.7
        for i,side in enumerate(("left","right")):
            x=MARGIN+i*(cw+gap)
            self.rect(slide,x,top+0.1,cw,chh,"surface",f"{side.title()} card",MSO_SHAPE.ROUNDED_RECTANGLE)
            hx=0
            if s[side].get("icon"):
                self.icon(slide,s[side]["icon"],x+0.35,top+0.48,0.62,f"{side.title()} icon",
                          "accent" if i==0 else "accent2")
                hx=0.8
            self.text(slide,x+0.35+hx,top+0.4,cw-0.7-hx,0.8,[s[side]["heading"]],30,f"{side.title()} heading",role="head",
                      color="accent" if i==0 else "accent2",bold=True,idx=idx)
            self.text(slide,x+0.35,top+1.35,cw-0.7,chh-1.5,["• "+p for p in s[side].get("points",[])],24,
                      f"{side.title()} points",gap=14,idx=idx)

    def l_quote(self,slide,s,idx):
        self.text(slide,1.3,0.4,1.8,2.1,["\u201c"],110,"Quote mark",role="head",color="accent",bold=True,idx=idx)
        self.text(slide,1.6,2.25,W-3.2,2.95,[s["quote"]],40,"Quote",role="head",anchor=MSO_ANCHOR.MIDDLE,idx=idx)
        if s.get("author"):
            self.rect(slide,1.7,5.45,0.9,0.05,"accent","Author rule")
            self.text(slide,2.75,5.2,W-4.4,0.6,[s["author"]],20,"Author",color="muted",idx=idx)

    def l_closing(self,slide,s,idx):
        self.text(slide,MARGIN,2.4,W-2*MARGIN,1.6,[s["title"]],48,"Title",role="head",bold=True,
                  align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.BOTTOM,idx=idx)
        self.rect(slide,W/2-1.0,4.15,2.0,0.08,"accent","Closing accent")
        if s.get("subtitle"):
            self.text(slide,MARGIN+1,4.45,W-2*MARGIN-2,1.2,[s["subtitle"]],24,"Subtitle",color="muted",
                      align=PP_ALIGN.CENTER,idx=idx)

    def build(self,spec,out):
        n=len(spec["slides"])
        for idx,s in enumerate(spec["slides"],1):
            slide=self.prs.slides.add_slide(self.blank)
            self.background(slide,idx,n)
            getattr(self,"l_"+s["layout"])(slide,s,idx)
            notes=[s["notes"]] if s.get("notes") else []
            notes+=[f"Image: {c['attribution']} ({c['landing']})" for c in self.credits if c["slide"]==idx]
            if notes:
                slide.notes_slide.notes_text_frame.text="\n".join(notes)
        if self.credits and spec.get("credits_slide",True):
            # Licensed photos are attributed on a closing credits slide (TASL).
            slide=self.prs.slides.add_slide(self.blank)
            self.background(slide,n+1,n+1)
            top=self.heading(slide,"Image credits",n+1)
            lines=[f"Slide {c['slide']}: {c['attribution']}" for c in self.credits]
            self.text(slide,MARGIN,top+0.1,W-2*MARGIN,H-top-0.6,lines,16,"Credits",color="muted",floor=12,idx=n+1)
            slide.notes_slide.notes_text_frame.text="[static] Attribution for licensed images."
        self.prs.save(out)


def _alpha(shape,percent):
    fill=shape.fill._xPr.find(qn("a:solidFill"))
    clr=fill[0]
    E.SubElement(clr,qn("a:alpha"),val=str(percent*1000))


def _arrow(connector):
    ln=connector.line._get_or_add_ln()
    E.SubElement(ln,qn("a:tailEnd"),type="triangle",w="med",len="med")


def _shape_text(b,shape,text,size,color):
    tf=shape.text_frame
    tf.word_wrap=True
    fam=b.fonts["head"]
    # An ellipse lays text out in its inscribed rectangle (~0.707 of each side).
    k=0.707 if shape.auto_shape_type==MSO_SHAPE.OVAL else 1.0
    w=shape.width/914400*k-0.2
    fitted,_=fit_size([text],fam,size,w+0.2,shape.height/914400*k-0.2,floor=10,bold=True)
    tf.text=text
    for p in tf.paragraphs:
        p.alignment=PP_ALIGN.CENTER
        for r in p.runs:
            r.font.size=Pt(fitted)
            r.font.bold=True
            r.font.name=fam
            r.font.color.rgb=b.rgb(color)


# ---------------------------------------------------------------------------
# Pipeline


def build(spec,out,work=None):
    """spec -> static deck -> Motion Director -> QA. Returns the verdict."""
    import forge_qa
    import motion_director as md
    errs=validate_spec(spec)
    if errs:
        return {"ok":False,"stage":"spec","errors":errs,"warnings":[]}
    out=Path(out)
    out.parent.mkdir(parents=True,exist_ok=True)
    static=Path(work or out.parent)/(out.stem+".static.pptx")
    b=Builder(spec)
    b.build(spec,static)
    model=md.deck_model(static)
    plan=md.draft(model,goal=spec.get("title",""),style=b.motion["style"],counters=bool(b.motion["counters"]),
                  morph=bool(b.motion["morph"]))
    report=md.apply(static,plan,out,force=True)  # the agent loop rebuilds the same output
    verdict=forge_qa.qa(out,b.theme)
    verdict["errors"]=[r for r in b.receipts if r.get("severity")=="error"]+verdict["errors"]
    verdict["warnings"]=[r for r in b.receipts if r.get("severity")!="error"]+verdict["warnings"]
    if not report["ok"]:
        verdict["errors"]+=[{"code":"MOTION_APPLY","slide":None,"object":None,"detail":p,"fix":"report this spec"}
                            for p in report["problems"]]
    verdict["warnings"]+=[{"code":"MOTION_WARNING","slide":None,"object":None,"detail":w,"fix":"check the preview"}
                          for w in report["warnings"]]
    verdict["ok"]=not verdict["errors"]
    verdict["generated_art"]=b.generated
    verdict["image_credits"]=b.credits
    verdict["icons"]=b.icons
    verdict["static_pptx"]=str(static)
    verdict["output"]=str(out)
    verdict["director_notes"]=plan["research_metadata"]["draft_notes"]
    return verdict


def render(pptx,outdir,seconds=8):
    """Render with desktop PowerPoint (Windows): MP4 via CreateVideo, the main
    sequence as PowerPoint parsed it, and one 6-frame strip per slide for the
    agent to look at. Returns paths; raises if PowerPoint/ffmpeg is missing."""
    import shutil
    import subprocess
    outdir=Path(outdir)
    outdir.mkdir(parents=True,exist_ok=True)
    if not (shutil.which("powershell") and shutil.which("ffmpeg")):
        raise RuntimeError("render needs Windows PowerShell + desktop PowerPoint + ffmpeg")
    ps1=Path(__file__).with_name("powerpoint_render.ps1")
    subprocess.run(["powershell","-NoProfile","-ExecutionPolicy","Bypass","-File",str(ps1),"-Pptx",str(pptx),
                    "-OutDir",str(outdir),"-SecondsPerSlide",str(seconds)],check=True,capture_output=True)
    video=outdir/"native_render.mp4"
    n=len(Presentation(str(pptx)).slides)
    dur=float(subprocess.run(["ffprobe","-v","error","-show_entries","format=duration","-of","csv=p=0",str(video)],
                             check=True,capture_output=True,text=True).stdout)
    per=dur/n
    sheets=[]
    for i in range(n):
        sheet=outdir/f"slide-{i+1:02d}.png"
        subprocess.run(["ffmpeg","-v","error","-y","-ss",f"{i*per:.3f}","-t",f"{per:.3f}","-i",str(video),
                        "-vf",f"fps={6/per:.4f},scale=480:-1,tile=6x1","-frames:v","1",str(sheet)],check=True)
        sheets.append(str(sheet))
    seq=json.loads((outdir/"native_sequence.json").read_text(encoding="utf-8-sig"))
    return {"ok":True,"stage":"render","powerpoint":f"{seq['powerpoint_version']} build {seq['build']}",
            "video":str(video),"sheets":sheets,"sequence":str(outdir/"native_sequence.json"),
            "note":"CreateVideo auto-advances clicks; it is PowerPoint's renderer, not an interactive slideshow"}


def main(argv=None):
    ap=argparse.ArgumentParser(description=__doc__,formatter_class=argparse.RawDescriptionHelpFormatter)
    sub=ap.add_subparsers(dest="cmd",required=True)
    sub.add_parser("schema")
    b=sub.add_parser("build")
    b.add_argument("spec",type=Path)
    b.add_argument("-o","--output",type=Path,required=True)
    q=sub.add_parser("qa")
    q.add_argument("pptx",type=Path)
    r=sub.add_parser("render")
    r.add_argument("pptx",type=Path)
    r.add_argument("-d","--dir",type=Path,required=True)
    r.add_argument("--seconds",type=int,default=8)
    a=ap.parse_args(argv)
    if a.cmd=="schema":
        result=schema()
    elif a.cmd=="build":
        result=build(json.loads(a.spec.read_text(encoding="utf-8")),a.output)
    elif a.cmd=="render":
        result=render(a.pptx,a.dir,a.seconds)
    else:
        import forge_qa
        result=forge_qa.qa(a.pptx)
    json.dump(result,sys.stdout,indent=1,ensure_ascii=False)
    print()
    return 0 if result.get("ok",True) else 1


if __name__=="__main__":
    sys.exit(main())
