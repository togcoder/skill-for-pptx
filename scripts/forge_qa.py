#!/usr/bin/env python3
"""Machine-readable QA for a finished deck (T025, Slide Forge).

Checks what an agent cannot see in XML alone and what a human reviewer would
flag: text that overflows its box, text over text, objects off the slide,
low contrast, tiny type, slides without motion and pictures that never move.
Every finding is {code, slide, object, detail, fix}; errors block delivery,
warnings are judgement calls. Works on any PPTX, not only Forge output.
"""
import sys
from pathlib import Path

from pptx import Presentation

sys.path.insert(0,str(Path(__file__).resolve().parent))
import pptx_animator as anim  # noqa: E402
from slide_forge import text_lines  # noqa: E402

NS=anim.NS
EMU_IN=914400


def _lum(hexcol):
    def ch(v):
        v/=255
        return v/12.92 if v<=0.03928 else ((v+0.055)/1.055)**2.4
    r,g,b=(int(hexcol[i:i+2],16) for i in (0,2,4))
    return 0.2126*ch(r)+0.7152*ch(g)+0.0722*ch(b)


def contrast(a,b):
    la,lb=sorted((_lum(a),_lum(b)),reverse=True)
    return (la+0.05)/(lb+0.05)


def _box(sh,sw,shh):
    if sh.left is None or sh.width is None:
        return None
    return (sh.left/sw,sh.top/shh,(sh.left+sh.width)/sw,(sh.top+sh.height)/shh)


def _area(b):
    return max(0,b[2]-b[0])*max(0,b[3]-b[1])


def _inter(a,b):
    return _area((max(a[0],b[0]),max(a[1],b[1]),min(a[2],b[2]),min(a[3],b[3])))


def _solid(xml_el):
    """Hex of the first direct solid fill (srgbClr) under an element, if any."""
    if xml_el is None:
        return None
    clr=xml_el.find("a:solidFill/a:srgbClr",NS)
    return clr.get("val").upper() if clr is not None else None


def _animated_spids(root):
    spids=set()
    timing=root.find("p:timing",NS)
    if timing is not None:
        spids={t.get("spid") for t in timing.iter(f"{{{anim.P}}}spTgt")}
    return spids


def qa(path,theme=None):
    prs=Presentation(str(path))
    sw,shh=prs.slide_width,prs.slide_height
    errors=[]
    warnings=[]
    metrics={"slides":len(prs.slides),"effects":0,"morph_transitions":0,"pictures":0,"pictures_moving":0,
             "slides_with_motion":0}

    def find(code,slide,obj,detail,fix,error=True):
        (errors if error else warnings).append({"code":code,"slide":slide,"object":obj,"detail":detail,"fix":fix})

    for idx,slide in enumerate(prs.slides,1):
        root=slide._element
        bg=_solid(root.find("p:cSld/p:bg/p:bgPr",NS)) or (theme or {}).get("bg","FFFFFF").upper()
        animated=_animated_spids(root)
        n_eff=len(root.findall(".//p:timing//p:cTn[@presetClass]",NS))
        metrics["effects"]+=n_eff
        has_morph=root.find(f".//{{{anim.P159}}}morph") is not None
        metrics["morph_transitions"]+=int(has_morph)
        if n_eff or has_morph:
            metrics["slides_with_motion"]+=1
        else:
            find("MOTION_NONE",idx,None,"slide has no animation and no Morph entry",
                 "give it a picture, accent or structured layout, or accept a static slide",error=False)
        shapes=[sh for sh in slide.shapes]
        boxes={sh.shape_id:_box(sh,sw,shh) for sh in shapes}
        fills=[(boxes[sh.shape_id],_solid(sh._element.find("p:spPr",NS))) for sh in shapes
               if boxes[sh.shape_id] and not sh.name.startswith("!!") and _solid(sh._element.find("p:spPr",NS))]
        texts=[]
        for sh in shapes:
            b=boxes[sh.shape_id]
            if b is None:
                continue
            name=sh.name
            if not name.startswith("!!") and (b[0]<-0.01 or b[1]<-0.01 or b[2]>1.01 or b[3]>1.01):
                find("OFF_SLIDE",idx,name,f"bbox {tuple(round(v,3) for v in b)}","move it inside the slide")
            if sh.shape_type==13:  # picture
                metrics["pictures"]+=1
                if _area(b)>=0.08:
                    moving=str(sh.shape_id) in animated
                    metrics["pictures_moving"]+=int(moving or has_morph)
                    if not moving and not has_morph:
                        find("PICTURE_FROZEN",idx,name,"large picture never moves",
                             "let the Director add ken-burns or give it a Morph partner",error=False)
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            if not name.startswith("__counter"):  # count-up proxies are stacked and shown one at a time
                texts.append((sh,b))
            runs=[r for p in sh.text_frame.paragraphs for r in p.runs]
            sizes=[r.font.size.pt for r in runs if r.font.size]
            if sizes and min(sizes)<12 and not name.startswith("Counter"):
                find("TYPE_TOO_SMALL",idx,name,f"{min(sizes):g} pt","keep body text >= 12 pt",error=False)
            # Overflow re-check with the same metric the builder uses.
            if sizes and sh.shape_type in (17,1):
                family=next((r.font.name for r in runs if r.font.name),"Segoe UI")
                paras=[p.text for p in sh.text_frame.paragraphs if p.text.strip()]
                lines=sum(text_lines(t,family,max(sizes),sh.width/EMU_IN-0.2) for t in paras)
                need=(lines*max(sizes)*1.18+6*(len(paras)-1))/72+0.1
                have=sh.height/EMU_IN
                if need>have*1.08:
                    find("TEXT_OVERFLOW",idx,name,f"needs ~{need:.2f} in, box {have:.2f} in","shorten or enlarge the box")
            # Contrast against the smallest filled shape underneath, else the slide.
            colors={r.font.color.rgb for r in runs if r.font.color and r.font.color.type is not None and r.font.color.rgb}
            under=[f for f in fills if f[0]!=b and _inter(f[0],b)>=0.9*_area(b)]
            back=min(under,key=lambda f:_area(f[0]))[1] if under else bg
            own=_solid(sh._element.find("p:spPr",NS))
            back=own or back
            for c in colors:
                ratio=contrast(str(c),back)
                need=3.0 if (sizes and max(sizes)>=24) else 4.5
                if ratio<need:
                    find("LOW_CONTRAST",idx,name,f"{c} on {back}: {ratio:.2f} < {need}","use the theme text color")
        for i,(a,ba) in enumerate(texts):
            for c,bc in texts[i+1:]:
                ov=_inter(ba,bc)
                if ov>0.15*min(_area(ba),_area(bc)):
                    find("TEXT_COLLISION",idx,f"{a.name} / {c.name}",f"overlap {ov:.3f} of slide","separate the boxes")
    metrics["motion_coverage"]=round(metrics["slides_with_motion"]/max(1,metrics["slides"]),2)
    return {"ok":not errors,"stage":"qa","errors":errors,"warnings":warnings,"metrics":metrics}
