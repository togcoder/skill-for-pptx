#!/usr/bin/env python3
"""Build the T019 synthetic "real-world report" fixture deck.

Development-only generator (needs python-pptx + Pillow). The committed fixture
at tests/fixtures/report_deck.pptx is what tests use; regenerating is optional.

The deck deliberately contains what ordinary business decks contain and the
earlier generated-deck fixtures did not: layout placeholders without slide-level
geometry, bullet bodies, KPI cards, a native chart, a process row with
connectors, a picture, a group, speaker notes, and one slide that already has
PowerPoint-authored paragraph animation. All data is synthetic.
"""
import argparse
import io
from pathlib import Path

from lxml import etree as E
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Emu, Inches, Pt

NAVY=RGBColor(0x1F,0x2A,0x44)
TEAL=RGBColor(0x0F,0x9D,0x8A)
SAND=RGBColor(0xF4,0xEF,0xE6)
GREY=RGBColor(0x5B,0x61,0x6E)

# PowerPoint-authored shape: "Fade, By paragraph, On click" on a body placeholder.
EXISTING_TIMING="""<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst><p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="4" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="5" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr><p:cTn id="6" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"><p:txEl><p:pRg st="0" end="0"/></p:txEl></p:spTgt></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set><p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="7" dur="500"/><p:tgtEl><p:spTgt spid="{spid}"><p:txEl><p:pRg st="0" end="0"/></p:txEl></p:spTgt></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par><p:par><p:cTn id="8" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="9" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:par><p:cTn id="10" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="clickEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:set><p:cBhvr><p:cTn id="11" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"><p:txEl><p:pRg st="1" end="1"/></p:txEl></p:spTgt></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set><p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="12" dur="500"/><p:tgtEl><p:spTgt spid="{spid}"><p:txEl><p:pRg st="1" end="1"/></p:txEl></p:spTgt></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst><p:bldP spid="{spid}" grpId="0" build="p"/></p:bldLst></p:timing>"""


def _title(slide,text):
    slide.shapes.title.text=text
    return slide.shapes.title


def _box(slide,x,y,w,h,text,size=14,bold=False,color=NAVY,name=None):
    tb=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=tb.text_frame
    tf.word_wrap=True
    tf.text=text
    run=tf.paragraphs[0].runs[0]
    run.font.size=Pt(size)
    run.font.bold=bold
    run.font.color.rgb=color
    if name:
        tb.name=name
    return tb


def _card(slide,x,y,w,h,fill=SAND,name=None):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb=fill
    shape.line.fill.background()
    if name:
        shape.name=name
    return shape


def _notes(slide,text):
    slide.notes_slide.notes_text_frame.text=text


def build(path):
    prs=Presentation()
    prs.slide_width=Inches(13.333)
    prs.slide_height=Inches(7.5)
    title_layout,content_layout,title_only=prs.slide_layouts[0],prs.slide_layouts[1],prs.slide_layouts[5]

    s=prs.slides.add_slide(title_layout)
    _title(s,"Q3 Operating Review")
    s.placeholders[1].text="Northwind Roasters (synthetic data) · October 2026"
    _notes(s,"Welcome. Today: how Q3 went, what drove it, and what we do next.")

    s=prs.slides.add_slide(content_layout)
    _title(s,"Agenda")
    body=s.placeholders[1].text_frame
    for i,line in enumerate(["Headline results","Revenue by channel","How we fixed fulfilment","Priorities for Q4"]):
        para=body.paragraphs[0] if i==0 else body.add_paragraph()
        para.text=line

    s=prs.slides.add_slide(title_only)
    _title(s,"Q3 at a glance")
    for i,(value,label) in enumerate([("$4.2M","Revenue, up from $3.1M"),("38%","Gross margin"),("1,250","New subscribers")]):
        x=0.9+i*4.0
        _card(s,x,2.2,3.5,3.0,name=f"KPI Card {i+1}")
        _box(s,x+0.3,2.6,2.9,1.2,value,size=44,bold=True,color=TEAL,name=f"KPI Value {i+1}")
        _box(s,x+0.3,3.9,2.9,0.8,label,size=16,color=GREY,name=f"KPI Label {i+1}")
    _notes(s,"Start with revenue: 4.2 million. Then margin held at 38 percent. Finally 1,250 new subscribers.")

    s=prs.slides.add_slide(title_only)
    _title(s,"Revenue by channel")
    data=CategoryChartData()
    data.categories=["Q4 25","Q1 26","Q2 26","Q3 26"]
    data.add_series("Retail",(1.4,1.5,1.7,2.0))
    data.add_series("Online",(1.0,1.2,1.4,2.2))
    frame=s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,Inches(0.8),Inches(1.7),Inches(7.8),Inches(5.0),data)
    frame.name="Channel Chart"
    chart=frame.chart
    chart.has_legend=True
    chart.legend.position=XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout=False
    _card(s,9.0,2.2,3.6,3.2,name="Takeaway Card")
    _box(s,9.3,2.5,3.0,2.6,"Online overtook retail for the first time in Q3.",size=20,bold=True,name="Takeaway")

    s=prs.slides.add_slide(title_only)
    _title(s,"How we fixed fulfilment")
    steps=["Forecast","Roast","Pack","Ship"]
    prev=None
    for i,label in enumerate(steps):
        x=0.8+i*3.2
        step=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(3.0),Inches(2.4),Inches(1.4))
        step.name=f"Step {i+1}"
        step.fill.solid(); step.fill.fore_color.rgb=NAVY if i<3 else TEAL
        step.line.fill.background()
        step.text_frame.text=label
        step.text_frame.paragraphs[0].runs[0].font.size=Pt(22)
        step.text_frame.paragraphs[0].runs[0].font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
        if prev is not None:
            conn=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,prev.left+prev.width,Inches(3.7),step.left,Inches(3.7))
            conn.name=f"Arrow {i}"
            conn.line.color.rgb=GREY
            conn.line.width=Pt(2)
        prev=step
    _box(s,0.8,5.0,11.8,0.8,"Lead time fell from 9 days to 4 days.",size=20,bold=True,color=TEAL,name="Result")
    # A decorative grouped badge (groups animate as one object).
    grp=s.shapes.add_group_shape()
    grp.name="Badge Group"
    a=grp.shapes.add_shape(MSO_SHAPE.OVAL,Inches(11.6),Inches(0.4),Inches(1.0),Inches(1.0))
    a.fill.solid(); a.fill.fore_color.rgb=TEAL; a.line.fill.background()
    grp.shapes.add_textbox(Inches(11.7),Inches(0.65),Inches(0.8),Inches(0.5)).text_frame.text="-55%"
    _notes(s,"Walk the four steps in order, then land the result: lead time down to four days.")

    s=prs.slides.add_slide(content_layout)
    _title(s,"Customer feedback")
    body=s.placeholders[1]
    body.text_frame.text="Delivery speed is the top compliment"
    body.text_frame.add_paragraph().text="Packaging waste is the top complaint"
    img=Image.new("RGB",(400,300),(244,239,230))
    d=ImageDraw.Draw(img)
    d.ellipse((120,70,280,230),fill=(15,157,138))
    buf=io.BytesIO(); img.save(buf,format="PNG"); buf.seek(0)
    body.left,body.top,body.width,body.height=Inches(0.8),Inches(1.8),Inches(7.0),Inches(4.8)
    pic=s.shapes.add_picture(buf,Inches(8.4),Inches(1.9),Inches(4.2),Inches(3.15))
    pic.name="Feedback Picture"
    timing=E.fromstring(EXISTING_TIMING.replace("{spid}",str(body.shape_id)))
    s._element.append(timing)
    _notes(s,"Existing animation: each feedback point fades in on click.")

    s=prs.slides.add_slide(content_layout)
    _title(s,"Priorities for Q4")
    body=s.placeholders[1].text_frame
    for i,line in enumerate(["Expand online subscriptions","Cut packaging waste by 30%","Open the second roastery"]):
        para=body.paragraphs[0] if i==0 else body.add_paragraph()
        para.text=line

    path.parent.mkdir(parents=True,exist_ok=True)
    prs.save(path)


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output",type=Path,nargs="?",default=Path(__file__).resolve().parents[1]/"tests"/"fixtures"/"report_deck.pptx")
    args=parser.parse_args()
    build(args.output)
    print(args.output)


if __name__=="__main__":
    main()
